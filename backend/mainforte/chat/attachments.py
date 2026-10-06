"""Attachment text extraction (plan §1, part 3): pulls searchable text out of an upload's raw
bytes at upload time, for types where that's possible. Kept separate from `routes.py` so the HTTP
handler stays about HTTP, not format-specific parsing.

Covers: plain text/markdown/CSV/JSON (decode as-is), PDF (`pypdf`), and the modern Office XML
formats — docx (`python-docx`), xlsx (`openpyxl`), pptx (`python-pptx`). Legacy binary Office
formats (`application/msword`, `application/vnd.ms-excel`, `application/vnd.ms-powerpoint`) are
accepted for upload (see `ALLOWED_PREFIXES`) but not extracted — their real content is a
proprietary binary layout with no lightweight pure-Python reader; add one only if a real need for
old `.doc`/`.xls`/`.ppt` shows up, rather than pulling in a fragile dependency speculatively.

Also builds Anthropic vision content blocks for image attachments -- images never reached the
model before this, same gap as document text."""
from __future__ import annotations

import base64
import io
import logging

log = logging.getLogger(__name__)

# Anthropic rejects an image source much past this; skip rather than send a request we know will
# be rejected (the model still sees the attachment reference via the message's text/name).
MAX_IMAGE_BYTES = 5 * 1024 * 1024

TEXT_PREFIXES = ("text/", "application/json")

OOXML_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OOXML_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
OOXML_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _extract_pdf(data: bytes) -> str | None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip() or None


def _extract_docx(data: bytes) -> str | None:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(p for p in parts if p.strip()) or None


def _extract_xlsx(data: bytes) -> str | None:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines).strip() or None


def _extract_pptx(data: bytes) -> str | None:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    lines: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        texts = [t for t in texts if t.strip()]
        if texts:
            lines.append(f"# Slide {i}")
            lines.extend(texts)
    return "\n".join(lines).strip() or None


_BINARY_EXTRACTORS = {
    "application/pdf": _extract_pdf,
    OOXML_DOCX: _extract_docx,
    OOXML_XLSX: _extract_xlsx,
    OOXML_PPTX: _extract_pptx,
}


def extract_text(*, content_type: str, data: bytes) -> str | None:
    """Best-effort text extraction for a stored upload's bytes. Returns None when the type isn't
    extractable or extraction fails — attachment text is a nice-to-have for search/recall, never a
    reason to fail the upload itself."""
    if content_type.startswith(TEXT_PREFIXES) or content_type == "text/csv":
        try:
            return data.decode("utf-8", errors="replace")
        except Exception:
            return None
    extractor = _BINARY_EXTRACTORS.get(content_type)
    if extractor is None:
        return None
    try:
        return extractor(data)
    except Exception:
        log.warning("attachment text extraction failed content_type=%s", content_type, exc_info=True)
        return None


def image_block(*, content_type: str, data: bytes) -> dict | None:
    """An Anthropic image content block for one attachment's bytes, or None if it's not an image
    or is too large to send (MAX_IMAGE_BYTES on the raw bytes -- base64 inflates further)."""
    if not content_type.startswith("image/") or len(data) > MAX_IMAGE_BYTES:
        return None
    return {
        "type": "image",
        "source": {"type": "base64", "media_type": content_type, "data": base64.b64encode(data).decode("ascii")},
    }
